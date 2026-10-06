"""Vuelca a texto (tmp_xl/*.txt) las fórmulas y valores de cada hoja del Excel oficial.
Sirve para explorar el libro y para generar los casos del oráculo (tools/gen_oracle_cases.py).
Uso: python tools/dump_workbook.py "<ruta al .xlsm>"
No modifica el archivo original.
"""
import os
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
path = sys.argv[1]
os.makedirs("tmp_xl", exist_ok=True)
wf = openpyxl.load_workbook(path, keep_vba=True, data_only=False)
wv = openpyxl.load_workbook(path, data_only=True)
for ws in wf.worksheets:
    v = wv[ws.title]
    with open(f"tmp_xl/{ws.title}.txt", "w", encoding="utf-8") as o:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                s = str(c.value)
                if s.startswith("="):
                    o.write(f"{c.coordinate}\t{s}\t=> {v[c.coordinate].value}\n")
                else:
                    o.write(f"{c.coordinate}\t{s}\n")
print("Hojas volcadas:", [ws.title for ws in wf.worksheets])
